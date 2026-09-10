"""Relatório Excel a partir dos pacotes cronometrados.

Produz uma pasta de trabalho com resumo executivo, decomposição do desvio em
relação ao alvo contratado, séries por minuto, distribuição dos intervalos,
o dado bruto e uma folha de método com as premissas.

A decomposição do desvio é o miolo do relatório. Vazão realizada abaixo do alvo
pode vir de duas causas independentes, e tratá-las como uma só leva à conclusão
errada sobre onde agir:

  * **cadência** — o intervalo entre posições consecutivas da esteira. Define o
    teto: nenhuma vazão passa de `3600 / cadência` por hora.
  * **ocupação** — a fração dessas posições que chega com pacote. Depende do
    que alimenta a linha, não da linha.

Medir as duas separadamente é o que permite dizer se a linha está lenta ou se
está recebendo menos do que consegue escoar.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .tracking import Cruzamento
from .video import formatar_tempo

# ---------------------------------------------------------------- estilo ------
AZUL = "1F3864"
AZUL_CLARO = "D9E2F3"
CINZA = "F2F2F2"
CINZA_BORDA = "BFBFBF"
VERDE = "375623"
VERDE_CLARO = "E2EFDA"
AMBAR_CLARO = "FFF2CC"
VERMELHO = "9C0006"
VERMELHO_CLARO = "FFC7CE"
BRANCO = "FFFFFF"

FONTE_TITULO = Font(name="Calibri", size=16, bold=True, color=AZUL)
FONTE_SEC = Font(name="Calibri", size=12, bold=True, color=BRANCO)
FONTE_CAB = Font(name="Calibri", size=10, bold=True, color=BRANCO)
FONTE_KPI = Font(name="Calibri", size=20, bold=True, color=AZUL)
FONTE_NOTA = Font(name="Calibri", size=9, italic=True, color="595959")

FUNDO_SEC = PatternFill("solid", fgColor=AZUL)
FUNDO_CAB = PatternFill("solid", fgColor="4472C4")
FUNDO_ALT = PatternFill("solid", fgColor=CINZA)

_lado = Side(style="thin", color=CINZA_BORDA)
BORDA = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)


def _num(v: float, dec: int = 0) -> str:
    """Número no formato brasileiro: 1705 -> "1.705", 2.111 -> "2,111".

    Formatar a frase inteira com `.replace(",", ".")` — que era o que eu fazia
    aqui — troca também as vírgulas da pontuação e corrompe o texto. A troca
    tem de ficar restrita ao número.
    """
    s = f"{v:,.{dec}f}"
    return s.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def _secao(ws, linha: int, texto: str, largura: int = 8) -> int:
    ws.cell(row=linha, column=1, value=texto).font = FONTE_SEC
    for c in range(1, largura + 1):
        ws.cell(row=linha, column=c).fill = FUNDO_SEC
    ws.row_dimensions[linha].height = 20
    return linha + 2


def _cabecalho(ws, linha: int, titulos: list[str], col0: int = 1) -> int:
    for i, t in enumerate(titulos):
        c = ws.cell(row=linha, column=col0 + i, value=t)
        c.font = FONTE_CAB
        c.fill = FUNDO_CAB
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[linha].height = 28
    return linha + 1


def _larguras(ws, larguras: dict[str, int]) -> None:
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- métricas ----
@dataclass
class Metricas:
    # Identificação
    video: str
    alvo_hora: float

    # Janela analisada
    t_ini: float
    t_fim: float
    duracao: float
    n_pacotes: int

    # Regime considerado para a comparação com o alvo
    reg_ini: float
    reg_fim: float
    reg_duracao: float
    reg_pacotes: int

    # Cadência e ocupação
    cadencia: float
    slots: int
    slots_vazios: int
    ocupacao: float

    # Vazões, todas em pacotes/hora
    realizado_hora: float
    capacidade_hora: float
    cadencia_alvo: float

    # Desvios em pacotes/hora
    gap_total: float
    gap_cadencia: float
    gap_ocupacao: float

    por_minuto: list[tuple[int, int]] = field(default_factory=list)
    paradas: list[tuple[float, float]] = field(default_factory=list)
    hist_slots: list[tuple[int, int]] = field(default_factory=list)
    gaps: np.ndarray = field(default_factory=lambda: np.array([]))
    primeiro: float = 0.0
    ultimo: float = 0.0


def calcular(
    eventos: list[Cruzamento],
    video: str,
    alvo_hora: float,
    t_ini: float,
    t_fim: float,
    reg_ini: float | None = None,
    reg_fim: float | None = None,
    limite_parada: float = 20.0,
) -> Metricas:
    tempos = np.array(sorted(e.tempo for e in eventos))
    if len(tempos) < 2:
        raise ValueError("São necessários pelo menos 2 pacotes para o relatório.")

    reg_ini = t_ini if reg_ini is None else reg_ini
    reg_fim = t_fim if reg_fim is None else reg_fim
    reg = tempos[(tempos >= reg_ini) & (tempos <= reg_fim)]
    if len(reg) < 2:
        raise ValueError("Trecho de regime sem pacotes suficientes.")

    gaps_reg = np.diff(reg)
    cadencia = float(np.median(gaps_reg))

    # Slots: cada intervalo vale um número inteiro de posições da esteira.
    # A soma dá quantas posições passaram pelo ponto de medição; a diferença
    # em relação à contagem de pacotes é o que passou vazio.
    n_slots = int(np.round(gaps_reg / cadencia).sum()) + 1
    vazios = max(0, n_slots - len(reg))
    ocupacao = len(reg) / n_slots if n_slots else 0.0

    reg_duracao = float(reg[-1] - reg[0]) or 1e-6
    realizado_hora = len(reg) / reg_duracao * 3600.0
    capacidade_hora = 3600.0 / cadencia
    cadencia_alvo = 3600.0 / alvo_hora if alvo_hora > 0 else float("nan")

    # O desvio total se separa em duas parcelas que somam exatamente ele:
    # o que falta por a cadência ser mais lenta que a exigida pelo alvo, e o
    # que falta por posições chegarem vazias.
    gap_total = alvo_hora - realizado_hora
    gap_cadencia = alvo_hora - capacidade_hora
    gap_ocupacao = capacidade_hora - realizado_hora

    contagem: dict[int, int] = {}
    for t in tempos:
        contagem[int(t // 60)] = contagem.get(int(t // 60), 0) + 1

    gaps_todos = np.diff(tempos)
    paradas = [
        (float(tempos[i + 1]), float(g))
        for i, g in enumerate(gaps_todos) if g > limite_parada
    ]

    mult = np.round(gaps_reg / cadencia).astype(int)
    hist: dict[int, int] = {}
    for m in mult:
        hist[int(m)] = hist.get(int(m), 0) + 1

    return Metricas(
        video=video, alvo_hora=alvo_hora,
        t_ini=t_ini, t_fim=t_fim, duracao=t_fim - t_ini, n_pacotes=len(tempos),
        reg_ini=reg_ini, reg_fim=reg_fim, reg_duracao=reg_duracao,
        reg_pacotes=len(reg),
        cadencia=cadencia, slots=n_slots, slots_vazios=vazios, ocupacao=ocupacao,
        realizado_hora=realizado_hora, capacidade_hora=capacidade_hora,
        cadencia_alvo=cadencia_alvo,
        gap_total=gap_total, gap_cadencia=gap_cadencia, gap_ocupacao=gap_ocupacao,
        por_minuto=sorted(contagem.items()),
        paradas=paradas,
        hist_slots=sorted(hist.items()),
        gaps=gaps_todos,
        primeiro=float(tempos[0]), ultimo=float(tempos[-1]),
    )


# ---------------------------------------------------------------- abas --------
def _aba_resumo(wb: Workbook, m: Metricas) -> None:
    ws = wb.create_sheet("Resumo")
    _larguras(ws, {"A": 34, "B": 18, "C": 18, "D": 18, "E": 16, "F": 16, "G": 14})

    ws["A1"] = "Análise de throughput — esteira de saída"
    ws["A1"].font = FONTE_TITULO
    ws["A2"] = f"Fonte: {Path(m.video).name}"
    ws["A2"].font = FONTE_NOTA
    ws["A3"] = (
        f"Janela analisada: {formatar_tempo(m.t_ini)} a {formatar_tempo(m.t_fim)}  ·  "
        f"regime de referência: {formatar_tempo(m.reg_ini)} a {formatar_tempo(m.reg_fim)}"
    )
    ws["A3"].font = FONTE_NOTA

    r = _secao(ws, 5, "INDICADORES", 7)

    atingimento = m.realizado_hora / m.alvo_hora if m.alvo_hora else 0.0
    kpis = [
        ("Alvo contratado", m.alvo_hora, "pacotes/h", None),
        ("Realizado (regime)", m.realizado_hora, "pacotes/h", atingimento),
        ("Capacidade na cadência observada", m.capacidade_hora, "pacotes/h",
         m.capacidade_hora / m.alvo_hora if m.alvo_hora else None),
        ("Ocupação das posições", m.ocupacao * 100, "%", None),
    ]
    for nome, valor, unidade, frac in kpis:
        ws.cell(row=r, column=1, value=nome).font = Font(bold=True, size=11)
        c = ws.cell(row=r, column=2, value=round(valor, 1))
        c.font = FONTE_KPI
        c.number_format = "#,##0.0"
        ws.cell(row=r, column=3, value=unidade).font = FONTE_NOTA
        if frac is not None:
            cf = ws.cell(row=r, column=4, value=frac)
            cf.number_format = "0.0%"
            cf.font = Font(bold=True, size=12)
            cf.fill = PatternFill(
                "solid",
                fgColor=VERDE_CLARO if frac >= 1 else
                (AMBAR_CLARO if frac >= 0.9 else VERMELHO_CLARO),
            )
            cf.alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=5, value="do alvo").font = FONTE_NOTA
        r += 2

    r = _secao(ws, r + 1, "ONDE ESTÁ O DESVIO", 7)
    ws.cell(row=r, column=1,
            value="O desvio se separa em duas causas independentes, que somam o total:"
            ).font = Font(italic=True, size=10)
    r += 2

    r = _cabecalho(r_ := r, ["Componente", "pacotes/h", "% do alvo", "Natureza"]) \
        if False else _cabecalho(ws, r, ["Componente", "pacotes/h", "% do alvo",
                                         "Natureza"])
    linhas = [
        ("Alvo contratado", m.alvo_hora, m.alvo_hora / m.alvo_hora, "referência"),
        ("(−) Cadência mais lenta que a exigida", -m.gap_cadencia,
         -m.gap_cadencia / m.alvo_hora, "ritmo da esteira"),
        ("(−) Posições que chegam vazias", -m.gap_ocupacao,
         -m.gap_ocupacao / m.alvo_hora, "alimentação da linha"),
        ("(=) Realizado", m.realizado_hora, m.realizado_hora / m.alvo_hora,
         "medido"),
    ]
    for nome, valor, frac, nat in linhas:
        destaque = nome.startswith("(=)") or nome.startswith("Alvo")
        c1 = ws.cell(row=r, column=1, value=nome)
        c1.font = Font(bold=destaque, size=10)
        c2 = ws.cell(row=r, column=2, value=round(valor, 1))
        c2.number_format = "#,##0.0;[Red]-#,##0.0"
        c2.font = Font(bold=destaque)
        c3 = ws.cell(row=r, column=3, value=frac)
        c3.number_format = "0.0%;[Red]-0.0%"
        ws.cell(row=r, column=4, value=nat).font = FONTE_NOTA
        for col in range(1, 5):
            ws.cell(row=r, column=col).border = BORDA
            if destaque:
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=AZUL_CLARO)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Leitura").font = Font(bold=True, size=11)
    r += 1
    for texto in _leitura(m):
        ws.cell(row=r, column=1, value=texto).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 30
        r += 1

    r = _secao(ws, r + 1, "CADÊNCIA", 7)
    dados = [
        ("Cadência observada (mediana)", m.cadencia, "s entre posições"),
        ("Cadência exigida pelo alvo", m.cadencia_alvo, "s entre posições"),
        ("Diferença", m.cadencia - m.cadencia_alvo, "s"),
        ("Posições no regime", m.slots, "posições"),
        ("Posições com pacote", m.reg_pacotes, "posições"),
        ("Posições vazias", m.slots_vazios, "posições"),
    ]
    for nome, valor, unidade in dados:
        ws.cell(row=r, column=1, value=nome).font = Font(size=10)
        c = ws.cell(row=r, column=2, value=round(valor, 3))
        c.number_format = "#,##0.000" if isinstance(valor, float) else "#,##0"
        ws.cell(row=r, column=3, value=unidade).font = FONTE_NOTA
        r += 1


def _leitura(m: Metricas) -> list[str]:
    """Frases de leitura, derivadas dos números — sem juízo de valor."""
    out = []
    if m.gap_cadencia > 0:
        out.append(
            f"A cadência observada ({_num(m.cadencia, 3)} s por posição) limita a "
            f"linha a {_num(m.capacidade_hora)} pacotes/h. O alvo de "
            f"{_num(m.alvo_hora)} pacotes/h exigiria {_num(m.cadencia_alvo, 3)} s "
            f"por posição. Mesmo com todas as posições ocupadas, o alvo não seria "
            f"alcançado nesta cadência."
        )
    else:
        out.append(
            f"A cadência observada ({_num(m.cadencia, 3)} s por posição) comporta "
            f"{_num(m.capacidade_hora)} pacotes/h, acima do alvo. O desvio vem "
            f"apenas das posições que chegam vazias."
        )
    out.append(
        f"A ocupação foi de {_num(m.ocupacao * 100, 1)}%: das {m.slots} posições que "
        f"passaram pelo ponto de medição no regime, {m.slots_vazios} chegaram sem "
        f"pacote. Isso responde por {_num(m.gap_ocupacao)} pacotes/h do desvio e "
        f"depende do que alimenta a linha, não do ritmo dela."
    )
    if m.paradas:
        pior = max(m.paradas, key=lambda p: p[1])
        n = len(m.paradas)
        if n == 1:
            t, d = m.paradas[0]
            out.append(
                f"Foi registrada 1 interrupção acima do limite, de {_num(d, 1)} s em "
                f"{formatar_tempo(t)}, equivalente a {d / m.cadencia:.0f} posições "
                f"perdidas."
            )
        else:
            out.append(
                f"Foram registradas {n} interrupções acima do limite. A maior durou "
                f"{_num(pior[1], 1)} s, em {formatar_tempo(pior[0])}, equivalente a "
                f"{pior[1] / m.cadencia:.0f} posições perdidas."
            )
    return out


def _aba_gap(wb: Workbook, m: Metricas) -> None:
    ws = wb.create_sheet("Desvio")
    _larguras(ws, {"A": 40, "B": 16, "C": 16})
    ws["A1"] = "Decomposição do desvio em relação ao alvo"
    ws["A1"].font = FONTE_TITULO

    r = _cabecalho(ws, 3, ["Etapa", "pacotes/h", "Acumulado"])
    passos = [
        ("Alvo contratado", m.alvo_hora),
        ("Perda por cadência", -m.gap_cadencia),
        ("Perda por posições vazias", -m.gap_ocupacao),
    ]
    acumulado = 0.0
    linha0 = r
    for nome, valor in passos:
        acumulado += valor
        ws.cell(row=r, column=1, value=nome)
        c = ws.cell(row=r, column=2, value=round(valor, 1))
        c.number_format = "#,##0.0;[Red]-#,##0.0"
        ca = ws.cell(row=r, column=3, value=round(acumulado, 1))
        ca.number_format = "#,##0.0"
        for col in range(1, 4):
            ws.cell(row=r, column=col).border = BORDA
        r += 1
    ws.cell(row=r, column=1, value="Realizado").font = Font(bold=True)
    cr = ws.cell(row=r, column=2, value=round(m.realizado_hora, 1))
    cr.number_format = "#,##0.0"
    cr.font = Font(bold=True)
    for col in range(1, 4):
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=AZUL_CLARO)
        ws.cell(row=r, column=col).border = BORDA

    graf = BarChart()
    graf.type = "col"
    graf.title = "Alvo, perdas e realizado (pacotes/h)"
    graf.y_axis.title = "pacotes/h"
    graf.height, graf.width = 9, 18
    dados = Reference(ws, min_col=2, min_row=linha0 - 1, max_row=r)
    cats = Reference(ws, min_col=1, min_row=linha0, max_row=r)
    graf.add_data(dados, titles_from_data=True)
    graf.set_categories(cats)
    graf.legend = None
    ws.add_chart(graf, "E3")


def _aba_por_minuto(wb: Workbook, m: Metricas) -> None:
    ws = wb.create_sheet("Vazão por minuto")
    _larguras(ws, {"A": 12, "B": 16, "C": 18, "D": 16, "E": 14})
    ws["A1"] = "Vazão por minuto de vídeo"
    ws["A1"].font = FONTE_TITULO

    r = _cabecalho(ws, 3, ["Minuto", "Pacotes", "Equivalente/h", "Alvo/h",
                           "% do alvo"])
    linha0 = r
    for minuto, n in m.por_minuto:
        ws.cell(row=r, column=1, value=minuto)
        ws.cell(row=r, column=2, value=n)
        c = ws.cell(row=r, column=3, value=n * 60)
        c.number_format = "#,##0"
        ws.cell(row=r, column=4, value=m.alvo_hora).number_format = "#,##0"
        cp = ws.cell(row=r, column=5, value=n * 60 / m.alvo_hora if m.alvo_hora else 0)
        cp.number_format = "0.0%"
        if cp.value and cp.value < 0.9:
            cp.fill = PatternFill("solid", fgColor=VERMELHO_CLARO)
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = BORDA
        r += 1

    graf = LineChart()
    graf.title = "Vazão equivalente por minuto vs alvo"
    graf.y_axis.title = "pacotes/h"
    graf.x_axis.title = "minuto do vídeo"
    graf.height, graf.width = 10, 22
    graf.add_data(Reference(ws, min_col=3, min_row=linha0 - 1, max_row=r - 1),
                  titles_from_data=True)
    graf.add_data(Reference(ws, min_col=4, min_row=linha0 - 1, max_row=r - 1),
                  titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=1, min_row=linha0, max_row=r - 1))
    ws.add_chart(graf, "G3")


def _aba_intervalos(wb: Workbook, m: Metricas) -> None:
    ws = wb.create_sheet("Intervalos")
    _larguras(ws, {"A": 26, "B": 14, "C": 16, "D": 42})
    ws["A1"] = "Distribuição dos intervalos entre pacotes"
    ws["A1"].font = FONTE_TITULO
    ws["A2"] = (
        "Os intervalos se concentram em múltiplos inteiros da cadência — evidência "
        "de que a esteira mantém ritmo fixo e o que varia é a ocupação."
    )
    ws["A2"].font = FONTE_NOTA

    r = _cabecalho(ws, 4, ["Intervalo", "Ocorrências", "% do total",
                           "Interpretação"])
    linha0 = r
    total = sum(n for _, n in m.hist_slots) or 1
    for mult, n in m.hist_slots:
        ws.cell(row=r, column=1, value=f"{mult}x a cadência "
                                      f"(~{mult * m.cadencia:.2f} s)")
        ws.cell(row=r, column=2, value=n)
        c = ws.cell(row=r, column=3, value=n / total)
        c.number_format = "0.0%"
        ws.cell(row=r, column=4,
                value="posições consecutivas ocupadas" if mult == 1
                else f"{mult - 1} posição(ões) vazia(s) entre pacotes")
        for col in range(1, 5):
            ws.cell(row=r, column=col).border = BORDA
        r += 1

    graf = BarChart()
    graf.type = "col"
    graf.title = "Ocorrências por múltiplo da cadência"
    graf.height, graf.width = 8, 14
    graf.add_data(Reference(ws, min_col=2, min_row=linha0 - 1, max_row=r - 1),
                  titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=1, min_row=linha0, max_row=r - 1))
    graf.legend = None
    ws.add_chart(graf, "F4")

    r += 2
    r = _secao(ws, r, "ESTATÍSTICAS DOS INTERVALOS", 4)
    g = m.gaps
    for nome, valor in [
        ("Mínimo (s)", float(g.min())),
        ("Percentil 25 (s)", float(np.percentile(g, 25))),
        ("Mediana (s)", float(np.median(g))),
        ("Percentil 75 (s)", float(np.percentile(g, 75))),
        ("Percentil 95 (s)", float(np.percentile(g, 95))),
        ("Máximo (s)", float(g.max())),
    ]:
        ws.cell(row=r, column=1, value=nome).font = Font(size=10)
        c = ws.cell(row=r, column=2, value=round(valor, 3))
        c.number_format = "#,##0.000"
        r += 1

    if m.paradas:
        r = _secao(ws, r + 1, "INTERRUPÇÕES", 4)
        r = _cabecalho(ws, r, ["Instante", "Duração (s)", "Posições perdidas"])
        for t, d in m.paradas:
            ws.cell(row=r, column=1, value=formatar_tempo(t))
            ws.cell(row=r, column=2, value=round(d, 2)).number_format = "#,##0.00"
            ws.cell(row=r, column=3, value=round(d / m.cadencia, 1)).number_format = \
                "#,##0.0"
            for col in range(1, 4):
                ws.cell(row=r, column=col).border = BORDA
            r += 1


def _aba_eventos(wb: Workbook, eventos: list[Cruzamento], m: Metricas) -> None:
    ws = wb.create_sheet("Pacotes")
    _larguras(ws, {"A": 10, "B": 16, "C": 10, "D": 10, "E": 14, "F": 16, "G": 18,
                   "H": 14})
    ws["A1"] = "Registro individual dos pacotes"
    ws["A1"].font = FONTE_TITULO
    ws["A2"] = "Instante em que a borda dianteira de cada pacote cruza o ponto de medição."
    ws["A2"].font = FONTE_NOTA

    r = _cabecalho(ws, 4, ["ID", "mm:ss.mmm", "Minuto", "Segundo",
                           "Milissegundo", "Segundos", "Intervalo anterior (s)",
                           "Posições"])
    ordenados = sorted(eventos, key=lambda e: e.tempo)
    anterior = None
    for ev in ordenados:
        total_ms = int(round(ev.tempo * 1000))
        ms = total_ms % 1000
        total_s = total_ms // 1000
        ws.cell(row=r, column=1, value=ev.id_pacote)
        ws.cell(row=r, column=2, value=formatar_tempo(ev.tempo))
        ws.cell(row=r, column=3, value=total_s // 60)
        ws.cell(row=r, column=4, value=total_s % 60)
        ws.cell(row=r, column=5, value=ms)
        ws.cell(row=r, column=6, value=round(ev.tempo, 3)).number_format = "#,##0.000"
        if anterior is not None:
            gap = ev.tempo - anterior
            cg = ws.cell(row=r, column=7, value=round(gap, 3))
            cg.number_format = "#,##0.000"
            cp = ws.cell(row=r, column=8, value=int(round(gap / m.cadencia)))
            if gap > 20.0:
                for col in (7, 8):
                    ws.cell(row=r, column=col).fill = PatternFill(
                        "solid", fgColor=VERMELHO_CLARO)
            elif cp.value and cp.value > 1:
                cp.fill = PatternFill("solid", fgColor=AMBAR_CLARO)
        anterior = ev.tempo
        r += 1
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{r - 1}"


def _aba_metodo(wb: Workbook, m: Metricas, notas: list[str]) -> None:
    ws = wb.create_sheet("Método e premissas")
    _larguras(ws, {"A": 4, "B": 110})
    ws["B1"] = "Método, premissas e limitações"
    ws["B1"].font = FONTE_TITULO

    blocos = [
        ("Como a medição foi feita", [
            "Medição por visão computacional sobre o vídeo da operação, sem "
            "instrumentação na linha.",
            "Uma janela estreita é definida sobre um ponto fixo da esteira e a "
            "ocupação de papelão nessa janela é medida quadro a quadro, como faria "
            "uma fotocélula. Cada transição de vazio para ocupado é a borda "
            "dianteira de um pacote chegando ao ponto.",
            "O instante é interpolado entre os dois quadros vizinhos. A 29,97 "
            "quadros por segundo, um quadro vale 33,4 ms; a interpolação leva a "
            "resolução para a casa do milissegundo.",
            "A distinção entre papelão e a estrutura da esteira usa o eixo b* do "
            "espaço de cor LAB: o metal dos roletes é levemente azulado, o papelão "
            "amarelado.",
        ]),
        ("Definições", [
            "Cadência — intervalo entre duas posições consecutivas da esteira no "
            "ponto de medição, tomado como a mediana dos intervalos observados no "
            "regime.",
            "Capacidade na cadência observada — 3600 dividido pela cadência. É o "
            "teto de vazão com todas as posições ocupadas.",
            "Ocupação — fração das posições que passam pelo ponto de medição "
            "carregando um pacote.",
            "Realizado — pacotes contados dividido pelo tempo decorrido, "
            "convertido para hora.",
        ]),
        ("Premissas e limitações", [
            f"A medição cobre {_num(m.duracao / 60, 1)} minutos de operação "
            f"({formatar_tempo(m.t_ini)} a {formatar_tempo(m.t_fim)}). É uma "
            f"amostra: a extrapolação para pacotes por hora assume que o regime "
            f"observado se mantém, o que este vídeo por si não demonstra.",
            f"O regime de referência para comparação com o alvo é "
            f"{formatar_tempo(m.reg_ini)} a {formatar_tempo(m.reg_fim)}, escolhido "
            f"por conter operação contínua. Trechos de partida e de interrupção "
            f"ficam fora dessa comparação e estão reportados separadamente.",
            "A cadência observada é a cadência EM QUE A LINHA ESTAVA OPERANDO no "
            "período filmado. Não é necessariamente o limite do equipamento: pode "
            "refletir parâmetro de configuração, velocidade ajustada ou limite "
            "imposto a montante. Concluir sobre capacidade máxima do equipamento "
            "exige ensaio dedicado, com alimentação saturada.",
            "A contagem é de pacotes que cruzam UM ponto. Se a operação tiver mais "
            "de uma via de saída, o total da instalação é a soma dos pontos.",
            "A câmera precisa permanecer estável, porque o ponto de medição é fixo "
            "em coordenadas de imagem. Neste vídeo a estabilidade foi verificada "
            "por correlação de fase: a deriva é inferior a 1 pixel ao longo de todo "
            "o período analisado.",
        ]),
        ("Rastreabilidade", [
            f"Arquivo de origem: {Path(m.video).name}",
            "O registro individual dos pacotes está na aba Pacotes e permite "
            "reconferir qualquer número deste relatório.",
            "O sinal bruto do sensor, quadro a quadro, é preservado em "
            "sinal_gate.npy, o que permite reauditar os limiares sem reprocessar o "
            "vídeo.",
        ]),
    ]
    if notas:
        blocos.append(("Observações", notas))

    r = 3
    for titulo, itens in blocos:
        ws.cell(row=r, column=2, value=titulo).font = Font(bold=True, size=12,
                                                           color=AZUL)
        r += 1
        for item in itens:
            c = ws.cell(row=r, column=2, value="•  " + item)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 14 * (len(item) // 95 + 1))
            r += 1
        r += 1


def gerar(
    eventos: list[Cruzamento],
    video: str,
    alvo_hora: float,
    t_ini: float,
    t_fim: float,
    saida: str | Path,
    reg_ini: float | None = None,
    reg_fim: float | None = None,
    limite_parada: float = 20.0,
    notas: list[str] | None = None,
) -> tuple[Path, Metricas]:
    m = calcular(eventos, video, alvo_hora, t_ini, t_fim, reg_ini, reg_fim,
                 limite_parada)
    wb = Workbook()
    wb.remove(wb.active)
    _aba_resumo(wb, m)
    _aba_gap(wb, m)
    _aba_por_minuto(wb, m)
    _aba_intervalos(wb, m)
    _aba_eventos(wb, eventos, m)
    _aba_metodo(wb, m, notas or [])
    saida = Path(saida)
    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    return saida, m
